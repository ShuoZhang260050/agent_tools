import io
import json
import sqlite3
import traceback
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import FastAPI, Depends, HTTPException, status, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response
from pydantic import BaseModel

from langchain_core.messages import HumanMessage

from agent.config import Settings
from agent.graph import build_graph_with_model
from agent.tools import get_tools
from agent.permissions import DEFAULT_PERMISSION, find_all_pending_approvals
from agent.auth import (
    create_access_token,
    get_current_user,
    decode_token,
)
from agent.memory.user_memory import (
    init_tables,
    register_user,
    authenticate,
    get_user_threads,
    get_user_threads_with_title,
    add_user_thread,
    update_thread_title,
    delete_user_thread,
    delete_thread_data,
    add_document,
    list_documents,
    delete_document,
    get_workspace,
    set_workspace,
)
from agent.sandbox.shadow import (
    get_active_workspace,
    clear_active_shadow,
    clear_all_user_shadows,
    ShadowManager,
    verify_shadow,
)
from agent.sandbox.snapshot import (
    init_snapshots_table,
    save_snapshot,
    restore_snapshot,
    get_latest_snapshot_id,
)

app = FastAPI(title="Agent API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
    allow_credentials=True,
)

_graphs: dict[str, object] = {}


def get_graph(model: str | None = None):
    """惰性构建并缓存 agent 图（按 model 缓存，共享 checkpointer）。"""
    settings = Settings()
    model = model or settings.llm_model
    if model not in _graphs:
        init_tables()
        _graphs[model] = build_graph_with_model(settings, model)
    return _graphs[model]


def _sse(obj) -> str:
    """将对象序列化为 SSE 事件字符串。"""
    return f"data: {json.dumps(obj, ensure_ascii=False)}\n\n"


def _serialize_msg(m) -> dict:
    """将 LangChain 消息对象序列化为前端可用的 dict。"""
    cls = type(m).__name__
    if cls == "HumanMessage":
        content = m.content
        if isinstance(content, list):
            text = ""
            image_url = None
            file_name = None
            file_text = None
            for part in content:
                if isinstance(part, dict):
                    if part.get("type") == "text":
                        part_text = part.get("text", "")
                        if part_text.startswith("[文件:"):
                            end = part_text.find("]")
                            if end > 0:
                                file_name = part_text[len("[文件: "):end].strip()
                                file_text = part_text[end + 1:].strip()
                                continue
                        text = (text + "\n" + part_text) if text else part_text
                    elif part.get("type") == "image_url":
                        image_url = part.get("image_url", {}).get("url")
            d = {"role": "user", "content": text.strip()}
            if image_url:
                d["image"] = image_url
            if file_name:
                d["file"] = file_name
                if file_text:
                    d["file_text"] = file_text
            return d
        return {"role": "user", "content": str(content)}
    if cls == "AIMessage":
        d = {"role": "assistant", "content": str(m.content or "")}
        if getattr(m, "tool_calls", None):
            d["tool_calls"] = [
                {"id": tc.get("id"), "name": tc.get("name"), "args": tc.get("args")}
                for tc in m.tool_calls
            ]
        return d
    if cls == "ToolMessage":
        return {
            "role": "tool",
            "content": str(m.content),
            "name": getattr(m, "name", ""),
            "tool_call_id": getattr(m, "tool_call_id", ""),
        }
    return {"role": "unknown", "content": str(getattr(m, "content", ""))}


class ChatReq(BaseModel):
    """流式对话请求模型。"""
    message: str
    thread_id: str | None = None
    image: str | None = None
    model: str | None = None
    attachment_text: str | None = None
    attachment_name: str | None = None
    permission: str = DEFAULT_PERMISSION


class ResumeReq(BaseModel):
    """恢复审批中断的请求模型。"""
    thread_id: str
    decision: str
    permission: str = DEFAULT_PERMISSION


class AuthReq(BaseModel):
    """用户认证请求模型。"""
    username: str
    password: str


@app.get("/health")
def health():
    """健康检查端点。"""
    return {"status": "ok"}


@app.get("/tools")
def tools():
    """列出所有已注册工具。"""
    return [{"name": t.name, "description": t.description} for t in get_tools()]


@app.get("/models")
def models():
    """列出可用模型及视觉能力。"""
    settings = Settings()
    avail = [m.strip() for m in settings.available_models.split(",") if m.strip()]
    if not avail:
        avail = [settings.llm_model]
    vision = {m.strip() for m in settings.vision_models.split(",") if m.strip()}
    return {"models": [{"name": m, "vision": m in vision} for m in avail]}


@app.get("/permissions")
def permissions():
    """列出权限选项及默认值。"""
    from agent.permissions import PERMISSION_CHOICES
    return {"choices": PERMISSION_CHOICES, "default": DEFAULT_PERMISSION}


@app.post("/register")
def register_endpoint(req: AuthReq):
    """用户注册端点。"""
    init_tables()
    try:
        user = register_user(req.username, req.password)
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=str(e))
    token = create_access_token(user["id"], user["username"])
    return {"token": token, "user": user}


@app.post("/login")
def login_endpoint(req: AuthReq):
    """用户登录端点。"""
    init_tables()
    user = authenticate(req.username, req.password)
    if not user:
        raise HTTPException(
            status_code=status.HTTP_401_UNAUTHORIZED,
            detail="用户名或密码错误",
        )
    token = create_access_token(user["id"], user["username"])
    return {"token": token, "user": user}


@app.get("/me")
def me(current: dict = Depends(get_current_user)):
    """获取当前登录用户信息。"""
    return current


@app.get("/sessions")
def list_sessions(current: dict = Depends(get_current_user)):
    """列出当前用户的会话（按最近活动排序）。"""
    settings = Settings()
    con = sqlite3.connect(settings.sqlite_path, check_same_thread=False)
    try:
        user_threads_map = {t["thread_id"]: t["title"] for t in get_user_threads_with_title(current["id"])}
        if not user_threads_map:
            return {"sessions": []}
        tids = list(user_threads_map.keys())
        placeholders = ",".join("?" for _ in tids)
        rows = con.execute(
            f"SELECT thread_id, MAX(checkpoint_id) AS latest "
            f"FROM checkpoints WHERE thread_id IN ({placeholders}) "
            f"GROUP BY thread_id ORDER BY latest DESC",
            tids,
        ).fetchall()
    finally:
        con.close()
    graph = get_graph()
    sessions = []
    for tid, latest in rows:
        try:
            state = graph.get_state({"configurable": {"thread_id": tid}})
            n = len(state.values.get("messages", []))
        except Exception:
            n = 0
        title = user_threads_map.get(tid) or tid[:8]
        sessions.append({"thread_id": tid, "title": title, "message_count": n, "updated_at": latest})
    return {"sessions": sessions}


@app.get("/sessions/{tid}")
def get_session(tid: str, current: dict = Depends(get_current_user)):
    """获取某会话的历史消息。"""
    user_threads = get_user_threads(current["id"])
    if tid not in user_threads:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="无权访问此会话")
    graph = get_graph()
    state = graph.get_state({"configurable": {"thread_id": tid}})
    msgs = state.values.get("messages", [])
    return {"thread_id": tid, "messages": [_serialize_msg(m) for m in msgs]}


def _stream_agent(graph, config, input_value, permission: str):
    """共享的 agent 流式生成器：处理 messages/updates 事件，结束后检测审批中断。

    input_value 可以是 {"messages": [...]} 或 langgraph Command(resume=...)。
    """
    try:
        for mode, data in graph.stream(
            input_value,
            config,
            stream_mode=["messages", "updates"],
        ):
            if mode == "messages":
                chunk, meta = data
                if meta.get("langgraph_node") == "model" and isinstance(chunk.content, str) and chunk.content:
                    yield _sse({"type": "token", "content": chunk.content})
            elif mode == "updates":
                for node, delta in data.items():
                    if not isinstance(delta, dict):
                        continue
                    for m in delta.get("messages", []):
                        if node == "model" and getattr(m, "tool_calls", None):
                            for tc in m.tool_calls:
                                yield _sse({
                                    "type": "tool_call",
                                    "id": tc.get("id"),
                                    "name": tc.get("name"),
                                    "args": tc.get("args"),
                                    "permission": permission,
                                })
                        elif node == "tools" and type(m).__name__ == "ToolMessage":
                            yield _sse({
                                "type": "tool_result",
                                "tool_call_id": getattr(m, "tool_call_id", None),
                                "name": getattr(m, "name", ""),
                                "content": str(m.content),
                            })
        pendings = find_all_pending_approvals(graph, config)
        if pendings:
            for p in pendings:
                yield _sse({"type": "approval_request", "permission": permission, **p})
        else:
            tid = config["configurable"]["thread_id"]
            uid = config["configurable"].get("user_id")
            shadow_path = get_active_workspace(uid, tid) if uid else None
            if shadow_path:
                real_ws = get_workspace(uid)
                if real_ws:
                    diff = ShadowManager.list_shadow_diff(shadow_path, real_ws)
                    has_changes = bool(diff["added"] or diff["modified"] or diff["deleted"])
                    yield _sse({
                        "type": "done",
                        "thread_id": tid,
                        "pending_sync": has_changes,
                        "diff": diff,
                    })
                else:
                    yield _sse({"type": "done", "thread_id": tid})
            else:
                yield _sse({"type": "done", "thread_id": tid})
    except Exception as e:
        tb = traceback.format_exc()
        try:
            with open("chat_errors.log", "a", encoding="utf-8") as f:
                f.write(f"[{datetime.now().isoformat()}] {tb}\n")
        except Exception:
            pass
        yield _sse({"type": "error", "message": f"{type(e).__name__}: {e}"})


@app.post("/chat")
async def chat(req: ChatReq, current: dict = Depends(get_current_user)):
    """流式对话端点，返回 SSE 事件流。"""
    graph = get_graph(req.model)
    tid = req.thread_id or str(uuid.uuid4())
    is_new = req.thread_id is None
    title = req.message[:30] if is_new else None
    add_user_thread(current["id"], tid, title=title)
    config = {"configurable": {"thread_id": tid, "user_id": current["id"],
                               "permission": req.permission},
              "metadata": {"thread_id": tid, "user_id": current["id"]}}
    callbacks = getattr(graph, "_tracing_callbacks", None)
    if callbacks:
        config["callbacks"] = list(callbacks)

    permission = req.permission
    image_data = req.image
    file_text = req.attachment_text
    file_name = req.attachment_name

    def gen():
        """生成 SSE 事件流的内部函数。"""
        try:
            if image_data:
                settings = Settings()
                vision = {m.strip() for m in settings.vision_models.split(",") if m.strip()}
                model_name = req.model or settings.llm_model
                if vision and model_name not in vision:
                    yield _sse({"type": "error", "message": f"模型 {model_name} 不支持图片，请切换到视觉模型后重试"})
                    yield _sse({"type": "done", "thread_id": tid})
                    return
                content = []
                if req.message:
                    content.append({"type": "text", "text": req.message})
                if file_text:
                    content.append({"type": "text", "text": f"[文件: {file_name or 'attachment'}]\n{file_text}"})
                content.append({"type": "image_url", "image_url": {"url": image_data}})
                msg = HumanMessage(content=content)
            elif file_text:
                content = []
                if req.message:
                    content.append({"type": "text", "text": req.message})
                content.append({"type": "text", "text": f"[文件: {file_name or 'attachment'}]\n{file_text}"})
                msg = HumanMessage(content=content)
            else:
                msg = HumanMessage(content=req.message)
            yield from _stream_agent(graph, config, {"messages": [msg]}, permission)
        except Exception as e:
            yield _sse({"type": "error", "message": f"{type(e).__name__}: {e}"})

    return StreamingResponse(
        gen(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/chat/resume")
def resume_chat(req: ResumeReq, current: dict = Depends(get_current_user)):
    """用户对挂起的工具调用给出「确认/取消」后，恢复 agent 执行。"""
    from langgraph.types import Command

    user_threads = get_user_threads(current["id"])
    if req.thread_id not in user_threads:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="无权访问此会话")
    graph = get_graph()
    config = {"configurable": {"thread_id": req.thread_id, "user_id": current["id"],
                               "permission": req.permission},
              "metadata": {"thread_id": req.thread_id, "user_id": current["id"]}}
    callbacks = getattr(graph, "_tracing_callbacks", None)
    if callbacks:
        config["callbacks"] = list(callbacks)

    pendings = find_all_pending_approvals(graph, config)
    if len(pendings) > 1:
        resume_value = Command(resume={
            p["interrupt_id"]: req.decision
            for p in pendings
            if p.get("interrupt_id")
        })
    else:
        resume_value = Command(resume=req.decision)

    return StreamingResponse(
        _stream_agent(graph, config, resume_value, req.permission),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.delete("/sessions/{tid}")
def delete_session(tid: str, current: dict = Depends(get_current_user)):
    """删除指定会话。"""
    user_threads = get_user_threads(current["id"])
    if tid not in user_threads:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="会话不存在")
    delete_user_thread(current["id"], tid)
    delete_thread_data(Settings().sqlite_path, tid)
    return {"status": "deleted"}


class RenameReq(BaseModel):
    """重命名会话的请求模型。"""
    title: str


@app.patch("/sessions/{tid}")
def rename_session(tid: str, req: RenameReq, current: dict = Depends(get_current_user)):
    """重命名会话。"""
    user_threads = get_user_threads(current["id"])
    if tid not in user_threads:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="会话不存在")
    update_thread_title(current["id"], tid, req.title)
    return {"thread_id": tid, "title": req.title}


@app.get("/documents")
def list_docs(current: dict = Depends(get_current_user)):
    """列出当前用户已上传的文档。"""
    return {"documents": list_documents(current["id"])}


@app.post("/documents")
async def upload_document(file: UploadFile = File(...), current: dict = Depends(get_current_user)):
    """上传文档到知识库（支持 .txt/.md/.pdf）。"""
    init_tables()
    filename = file.filename or "unknown.txt"
    suffix = Path(filename).suffix.lower()
    raw = await file.read()
    if suffix == ".pdf":
        try:
            import pymupdf
        except ImportError:
            raise HTTPException(status_code=400, detail="PDF 支持未启用：未安装 pymupdf")
        try:
            doc = pymupdf.open(stream=raw, filetype="pdf")
            text = "\n".join(page.get_text("text") for page in doc)
            doc.close()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"PDF 解析失败: {e}")
    elif suffix in (".txt", ".md", ".markdown"):
        text = raw.decode("utf-8", errors="replace")
    elif suffix in (".doc", ".docx"):
        try:
            import docx
        except ImportError:
            raise HTTPException(status_code=400, detail="Word 支持未启用：未安装 python-docx")
        try:
            doc = docx.Document(io.BytesIO(raw))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Word 解析失败: {e}")
    else:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {suffix}（支持 .txt/.md/.pdf/.docx）")
    if not text.strip():
        raise HTTPException(status_code=400, detail="文件内容为空")

    from agent.memory.vectorstore import build_embeddings, ingest_document
    settings = Settings()
    embeddings = build_embeddings(settings)
    doc_id = add_document(current["id"], filename, 0)
    try:
        chunk_count = ingest_document(
            current["id"], doc_id, text, filename, embeddings,
            settings.rag_chunk_size, settings.rag_chunk_overlap,
        )
    except Exception as e:
        delete_document(current["id"], doc_id)
        raise HTTPException(status_code=500, detail=f"文档向量化失败: {e}")
    return {"doc_id": doc_id, "filename": filename, "chunks": chunk_count}


@app.delete("/documents/{doc_id}")
def delete_doc(doc_id: int, current: dict = Depends(get_current_user)):
    """删除指定文档及其向量。"""
    docs = list_documents(current["id"])
    if not any(d["id"] == doc_id for d in docs):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="文档不存在")
    delete_document(current["id"], doc_id)
    return {"status": "deleted"}


@app.get("/traces")
def get_traces_api(thread_id: str | None = None, limit: int = 50, current: dict = Depends(get_current_user)):
    """获取调用追踪记录。"""
    from agent.memory.tracing import get_traces
    return {"traces": get_traces(thread_id=thread_id, limit=limit)}


@app.post("/extract-text")
async def extract_text(file: UploadFile = File(...), current: dict = Depends(get_current_user)):
    """提取文件文本（用于附加到对话，不入知识库）。"""
    filename = file.filename or "unknown.txt"
    suffix = Path(filename).suffix.lower()
    raw = await file.read()
    if len(raw) > 10 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="文件超过 10MB，建议上传到知识库")
    if suffix == ".pdf":
        try:
            import pymupdf
        except ImportError:
            raise HTTPException(status_code=400, detail="PDF 支持未启用：未安装 pymupdf")
        try:
            doc = pymupdf.open(stream=raw, filetype="pdf")
            text = "\n".join(page.get_text("text") for page in doc)
            doc.close()
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"PDF 解析失败: {e}")
    elif suffix in (".txt", ".md", ".markdown", ".csv", ".json", ".log"):
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("gbk", errors="replace")
    else:
        try:
            text = raw.decode("utf-8")
        except UnicodeDecodeError:
            text = raw.decode("gbk", errors="replace")
    if not text.strip():
        raise HTTPException(status_code=400, detail="文件内容为空")
    return {"text": text, "filename": filename}


@app.get("/workspace")
def get_workspace_api(current: dict = Depends(get_current_user)):
    """获取当前用户的工作空间路径。"""
    return {"workspace": get_workspace(current["id"])}


@app.post("/workspace")
def set_workspace_api(path: str, current: dict = Depends(get_current_user)):
    """设置当前用户的工作空间路径。"""
    p = Path(path)
    if not p.exists():
        raise HTTPException(status_code=400, detail="路径不存在")
    if not p.is_dir():
        raise HTTPException(status_code=400, detail="路径不是目录")
    saved = set_workspace(current["id"], str(p.resolve()))
    clear_all_user_shadows(current["id"])
    return {"workspace": saved}


@app.get("/workspace/browse")
def browse_workspace_api(path: str = ".", current: dict = Depends(get_current_user)):
    """列出目录内容，辅助用户选择工作空间。"""
    p = Path(path).resolve()
    if not p.is_dir():
        raise HTTPException(status_code=400, detail="不是目录")
    dirs = []
    try:
        for entry in sorted(p.iterdir(), key=lambda e: e.name.lower()):
            if entry.is_dir() and not entry.name.startswith("."):
                dirs.append({"name": entry.name, "path": str(entry)})
    except PermissionError:
        raise HTTPException(status_code=400, detail="无权限访问该目录")
    return {"current": str(p), "parent": str(p.parent) if str(p.parent) != str(p) else None, "dirs": dirs[:100]}


@app.get("/workspace/drives")
def list_drives_api(current: dict = Depends(get_current_user)):
    """列出系统可用的盘符。"""
    import string
    drives = []
    for letter in string.ascii_uppercase:
        drive = f"{letter}:\\"
        if Path(drive).exists():
            drives.append(drive)
    if not drives:
        drives = ["/"]
    return {"drives": drives}


@app.get("/workspace/diff")
def get_workspace_diff_api(thread_id: str, current: dict = Depends(get_current_user)):
    """获取当前 shadow 与 real workspace 的差异。"""
    uid = current["id"]
    shadow_path = get_active_workspace(uid, thread_id)
    if not shadow_path:
        return {"diff": {"added": [], "modified": [], "deleted": []}, "pending_sync": False}
    real_ws = get_workspace(uid)
    if not real_ws:
        return {"diff": {"added": [], "modified": [], "deleted": []}, "pending_sync": False}
    diff = ShadowManager.list_shadow_diff(shadow_path, real_ws)
    has_changes = bool(diff["added"] or diff["modified"] or diff["deleted"])
    return {"diff": diff, "pending_sync": has_changes}


@app.post("/workspace/sync")
def sync_workspace_api(thread_id: str, verify_command: str = "",
                        current: dict = Depends(get_current_user)):
    """将 shadow 工作空间的变更同步到真实工作空间。同步前自动快照。
    若提供 verify_command，则先在 shadow 中运行验证，失败则拒绝同步。
    """
    uid = current["id"]
    shadow_path = get_active_workspace(uid, thread_id)
    if not shadow_path:
        raise HTTPException(status_code=400, detail="没有活跃的 shadow 工作空间")
    real_ws = get_workspace(uid)
    if not real_ws:
        raise HTTPException(status_code=400, detail="未设置真实工作空间")

    verify_result = None
    if verify_command:
        verify_result = verify_shadow(shadow_path, verify_command)
        if not verify_result["passed"]:
            return {
                "synced": 0,
                "verified": False,
                "verify_output": verify_result["output"],
                "message": "验证失败，已阻止同步。请修复后重试。",
            }

    diff = ShadowManager.list_shadow_diff(shadow_path, real_ws)
    if not (diff["added"] or diff["modified"] or diff["deleted"]):
        clear_active_shadow(uid, thread_id)
        return {"synced": 0, "verified": True, "message": "无变更需要同步"}

    snapshot_id = save_snapshot(Settings().sqlite_path, uid, thread_id, real_ws, diff)
    result = ShadowManager.apply_shadow_to_real(shadow_path, real_ws)
    clear_active_shadow(uid, thread_id)
    return {
        "synced": result["synced"],
        "bytes": result["bytes"],
        "snapshot_id": snapshot_id,
        "verified": verify_result is None or verify_result["passed"],
    }


@app.post("/workspace/verify")
def verify_workspace_api(thread_id: str, command: str,
                         current: dict = Depends(get_current_user)):
    """在 shadow 工作空间中运行验证命令（如 pytest、ruff check），返回结果。"""
    uid = current["id"]
    shadow_path = get_active_workspace(uid, thread_id)
    if not shadow_path:
        raise HTTPException(status_code=400, detail="没有活跃的 shadow 工作空间")
    result = verify_shadow(shadow_path, command)
    return result


@app.post("/workspace/revert")
def revert_workspace_api(thread_id: str, current: dict = Depends(get_current_user)):
    """从最新快照恢复真实工作空间。"""
    uid = current["id"]
    snapshot_id = get_latest_snapshot_id(Settings().sqlite_path, uid, thread_id)
    if not snapshot_id:
        raise HTTPException(status_code=404, detail="没有可恢复的快照")
    result = restore_snapshot(Settings().sqlite_path, snapshot_id)
    clear_active_shadow(uid, thread_id)
    return result


@app.get("/screenshots/{filename}")
def get_screenshot(filename: str):
    """提供浏览器截图文件。"""
    from pathlib import Path
    path = Path(Settings().sqlite_path).parent / "screenshots" / filename
    if not path.exists():
        raise HTTPException(status_code=404, detail="截图不存在")
    return Response(content=path.read_bytes(), media_type="image/png")
